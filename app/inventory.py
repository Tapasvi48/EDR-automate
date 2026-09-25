"""LOB inventory: file parsing, template column mapping, versioning/diffing and EDR verification."""
import csv
import hashlib
import io
import json
import re
import uuid
from collections import OrderedDict
from datetime import datetime, timedelta, timezone

from . import config, db

FIELDS = config.INVENTORY_FIELD_KEYS
ALIASES = {
    "ip": ["ip", "ipaddress", "ipaddr", "hostip", "serverip", "primaryip", "managementip", "nodeip", "privateip", "ipv4"],
    "node_name": ["nodename", "hostname", "host", "servername", "computername", "devicename", "assetname", "machinename", "name", "node", "server"],
    "msp": ["msp", "mspname", "managedby", "managedserviceprovider", "serviceprovider", "vendor", "supportvendor", "partner", "supportpartner"],
    "node_type": ["nodetype", "type", "assettype", "devicetype", "servertype", "role", "category", "hosttype"],
    "domain": ["domain", "addomain", "domainname", "workgroup"],
    "live": ["livenonlive", "live", "livestatus", "livenonlivestatus", "status", "environment", "prodstatus"],
    "os": ["os", "operatingsystem", "osname", "osversion", "platform", "osdetails"],
    "edr_feasible": ["edrfeasible", "feasible", "edrfeasibility", "feasibility", "edrapplicable", "applicable"],
    "edr_installed": ["edrinstalled", "edr", "edrstatus", "crowdstrikeinstalled", "agentinstalled", "csinstalled", "falconinstalled", "installed"],
    "remarks": ["remarks", "remark", "comments", "comment", "notes", "note", "justification"],
}
YES = {"yes", "y", "true", "1", "installed", "done", "present", "feasible", "ok", "available", "deployed"}
NO = {"no", "n", "false", "0", "notinstalled", "not installed", "pending", "notfeasible", "not feasible", "missing", "absent", "not deployed"}
LIVE = {"live", "yes", "y", "prod", "production", "active", "running", "inuse", "in use"}
NONLIVE = {"nonlive", "non live", "non-live", "notlive", "not live", "no", "n", "decom", "decommissioned",
           "inactive", "retired", "shutdown", "poweredoff", "powered off", "offline"}
KEY_FIELDS = {"ip": "IP", "node_name": "Node Name", "ip_node_name": "IP + Node Name"}

_parsed_cache = OrderedDict()


def _hnorm(h):
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    return str(v).strip()


def norm_yes_no(v):
    s = str(v or "").strip()
    low = s.lower()
    if not low:
        return ""
    if low in YES:
        return "Yes"
    if low in NO:
        return "No"
    return s


def norm_live(v):
    s = str(v or "").strip()
    low = s.lower()
    if not low:
        return ""
    if low in LIVE:
        return "Live"
    if low in NONLIVE or low.replace(" ", "") in NONLIVE:
        return "Non Live"
    return s


def norm_ip(v):
    s = str(v or "").strip()
    # cells sometimes hold "10.1.1.1, 10.1.1.2" or "10.1.1.1 / eth0" -> first IP
    m = re.search(r"\b\d{1,3}(?:\.\d{1,3}){3}\b", s)
    return m.group(0) if m else s


# ------------------------------------------------------------------ parsing
def save_upload(filename, data: bytes):
    token = uuid.uuid4().hex
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", filename or "upload")
    path = config.UPLOAD_DIR / f"{token}__{safe}"
    path.write_bytes(data)
    return token


def _upload_path(token):
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        raise ValueError("Invalid upload token")
    for p in config.UPLOAD_DIR.glob(f"{token}__*"):
        return p
    raise ValueError("Upload expired - please upload the file again")


def parse_upload(token, sheet=None, header_row=None):
    """Returns dict(filename, sheets, sheet, header_row, headers, rows)."""
    ck = (token, sheet, header_row)
    if ck in _parsed_cache:
        return _parsed_cache[ck]
    path = _upload_path(token)
    filename = path.name.split("__", 1)[1]
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        ws = wb[sheet] if sheet in sheets else wb[sheets[0]]
        sheet = ws.title
        raw = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif ext in (".csv", ".txt"):
        text = path.read_bytes().decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        raw = [[_cell(v) for v in r] for r in csv.reader(io.StringIO(text), dialect)]
        sheets, sheet = ["CSV"], "CSV"
    else:
        raise ValueError("Unsupported file type. Upload .xlsx, .xlsm or .csv (save .xls as .xlsx first).")

    if header_row is None:
        header_row = 1
        for i, r in enumerate(raw[:30]):
            if sum(1 for v in r if v) >= 2:
                header_row = i + 1
                break
    header_row = max(1, int(header_row))
    headers = raw[header_row - 1] if len(raw) >= header_row else []
    # de-duplicate / fill blank headers
    seen, hdrs = {}, []
    for i, h in enumerate(headers):
        h = h or f"Column {i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        hdrs.append(h)
    body = [r for r in raw[header_row:] if any(v for v in r)]
    body = [(r + [""] * len(hdrs))[:len(hdrs)] for r in body]
    out = {"filename": filename, "sheets": sheets, "sheet": sheet, "header_row": header_row, "headers": hdrs, "rows": body}
    _parsed_cache[ck] = out
    while len(_parsed_cache) > 8:
        _parsed_cache.popitem(last=False)
    return out


def suggest_mapping(headers, template=None):
    mapping, used = {}, set()
    norm = {h: _hnorm(h) for h in headers}
    if template:
        tmap = db.jloads(template["mapping"], {})
        for f, src in tmap.items():
            for h in headers:
                if src and norm[h] == _hnorm(src) and h not in used:
                    mapping[f] = h
                    used.add(h)
                    break
    for f in FIELDS:
        if f in mapping:
            continue
        for alias in ALIASES[f]:
            hit = next((h for h in headers if h not in used and norm[h] == alias), None)
            if hit:
                mapping[f] = hit
                used.add(hit)
                break
    for f in FIELDS:  # looser "contains" pass
        if f in mapping:
            continue
        for alias in ALIASES[f][:3]:
            hit = next((h for h in headers if h not in used and len(alias) > 3 and alias in norm[h]), None)
            if hit:
                mapping[f] = hit
                used.add(hit)
                break
    return mapping


def make_key(item, key_field):
    ip = (item.get("ip") or "").lower()
    nn = db.norm_hostname(item.get("node_name"))
    if key_field == "node_name":
        base = nn or ip
    elif key_field == "ip_node_name":
        base = f"{ip}|{nn}" if (ip or nn) else ""
    else:
        base = ip or nn
    # the same IP may legitimately be listed by two MSPs of one LOB -> MSP is part of the identity
    msp = (item.get("msp") or "").strip().lower()
    return f"{msp}|{base}" if (msp and base) else base


def build_items(parsed, mapping, key_field, forced_msp=None, forced_type=None, key_prefix=""):
    headers = parsed["headers"]
    idx = {h: i for i, h in enumerate(headers)}
    mapped_headers = set(v for v in mapping.values() if v)
    extra_headers = [h for h in headers if h not in mapped_headers]
    items, warnings, dups = OrderedDict(), [], 0
    no_key = 0
    for n, r in enumerate(parsed["rows"], start=parsed["header_row"] + 1):
        it = {}
        for f in FIELDS:
            h = mapping.get(f)
            it[f] = r[idx[h]] if h in idx else ""
        it["ip"] = norm_ip(it["ip"])
        it["msp"] = (forced_msp or it["msp"] or "").strip()
        if forced_type:
            it["node_type"] = forced_type
        it["live"] = norm_live(it["live"])
        it["edr_feasible"] = norm_yes_no(it["edr_feasible"])
        it["edr_installed"] = norm_yes_no(it["edr_installed"])
        it["extra"] = {h: r[idx[h]] for h in extra_headers if r[idx[h]]}
        key = make_key(it, key_field)
        key = key_prefix + key if key else key
        if not key:
            no_key += 1
            continue
        if key in items:
            # same key repeated in the file: keep the first row, count the repeats and tag the item as duplicate
            dups += 1
            items[key]["file_dups"] = items[key].get("file_dups", 0) + 1
            if dups <= 25:
                warnings.append(f"Row {n}: duplicate key '{key}' - merged into the first occurrence (row {items[key]['_row']})")
            continue
        it["_row"] = n
        it["file_dups"] = 0
        items[key] = it
    for it in items.values():
        it.pop("_row", None)
    if dups > 25:
        warnings.append(f"... {dups - 25} more duplicate rows merged")
    if no_key:
        warnings.append(f"{no_key} rows skipped: no IP / Node Name value")
    return items, warnings


def merge_scope(prev_items, new_items, scope_msp):
    """MSP-scoped upload: keep every other MSP's rows from the previous version, replace this MSP's rows."""
    scope = (scope_msp or "").strip().lower()
    merged = OrderedDict((k, v) for k, v in prev_items.items() if (v.get("msp") or "").strip().lower() != scope)
    merged.update(new_items)
    return merged


def _hash(it):
    payload = json.dumps([it.get(f, "") for f in FIELDS] + [it.get("extra") or {}], sort_keys=True)
    return hashlib.sha1(payload.encode()).hexdigest()


def diff_items(old, new):
    """old/new: {key: item}. Returns (added_keys, removed_keys, modified {key: [(field, old, new)]}, unchanged_count)."""
    added = [k for k in new if k not in old]
    removed = [k for k in old if k not in new]
    modified, unchanged = {}, 0
    for k in new:
        if k not in old:
            continue
        a, b = old[k], new[k]
        if _hash(a) == _hash(b):
            unchanged += 1
            continue
        ch = [(f, a.get(f) or "", b.get(f) or "") for f in FIELDS if (a.get(f) or "") != (b.get(f) or "")]
        ea, eb = a.get("extra") or {}, b.get("extra") or {}
        for col in sorted(set(ea) | set(eb)):
            if (ea.get(col) or "") != (eb.get(col) or ""):
                ch.append((col, ea.get(col) or "", eb.get(col) or ""))
        if ch:
            modified[k] = ch
        else:
            unchanged += 1
    return added, removed, modified, unchanged


def load_version_items(c, version_id):
    out = OrderedDict()
    if not version_id:
        return out
    for r in c.execute("SELECT * FROM inventory_rows WHERE version_id=? ORDER BY rowid", (version_id,)):
        d = dict(r)
        d["extra"] = db.jloads(d["extra"], {})
        out[d["item_key"]] = d
    return out


def _diff_summary(items_old, items_new, sample=200):
    added, removed, modified, unchanged = diff_items(items_old, items_new)
    fld_label = dict(config.INVENTORY_FIELDS)
    return {
        "added": len(added), "removed": len(removed), "modified": len(modified), "unchanged": unchanged,
        "total": len(items_new),
        "samples": {
            "added": [{"key": k, **_brief(items_new[k])} for k in added[:sample]],
            "removed": [{"key": k, **_brief(items_old[k])} for k in removed[:sample]],
            "modified": [{"key": k, "node_name": items_new[k].get("node_name"), "ip": items_new[k].get("ip"),
                          "changes": [{"field": fld_label.get(f, f), "old": o, "new": n} for f, o, n in ch]}
                         for k, ch in list(modified.items())[:sample]],
        },
    }, (added, removed, modified, unchanged)


def _brief(it):
    return {f: it.get(f, "") for f in FIELDS}


# ------------------------------------------------------------------ commit / versions
ROW_COLS = ", ".join(FIELDS)


# ------------------------------------------------------------------ inventory streams (main + one per type)
def get_type(c, lob_id, type_id):
    if not type_id:
        return None
    t = db.one(c, "SELECT * FROM lob_types WHERE id=? AND lob_id=?", (type_id, lob_id))
    if not t:
        raise ValueError("Inventory type not found in this LOB")
    return t


def stream_version_id(c, lob, type_id):
    """Current version of the main inventory (type_id None) or of one inventory type."""
    if not type_id:
        return lob["current_version_id"]
    t = get_type(c, lob["id"], type_id)
    return t["current_version_id"]


def type_key_prefix(type_id):
    # keeps item keys of different inventory types apart in inventory_current
    return f"t{type_id}|" if type_id else ""


def commit_version(c, lob_id, items, *, filename, note, uploaded_by, template_id, key_field, mapping,
                   warnings=None, restored_from=None, scope_msp=None, type_id=None):
    lob = db.one(c, "SELECT * FROM lobs WHERE id=?", (lob_id,))
    if not lob:
        raise ValueError("LOB not found")
    type_id = type_id or None
    prev_items = load_version_items(c, stream_version_id(c, lob, type_id))
    summary, (added, removed, modified, unchanged) = _diff_summary(prev_items, items, sample=0)
    vno = (c.execute("SELECT MAX(version_no) FROM inventory_versions WHERE lob_id=? AND type_id IS ?",
                     (lob_id, type_id)).fetchone()[0] or 0) + 1
    vid = c.execute(
        """INSERT INTO inventory_versions(lob_id, version_no, filename, note, uploaded_by, uploaded_at, template_id,
           key_field, mapping, row_count, added, removed, modified, unchanged, restored_from, scope_msp, warnings, type_id)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (lob_id, vno, filename, note, uploaded_by, db.now_iso(), template_id, key_field, json.dumps(mapping),
         len(items), len(added), len(removed), len(modified), unchanged, restored_from, scope_msp, json.dumps(warnings or []),
         type_id)).lastrowid
    ph = ",".join("?" * len(FIELDS))
    c.executemany(
        f"""INSERT INTO inventory_rows(version_id, item_key, {ROW_COLS}, extra, row_hash, file_dups) VALUES (?,?,{ph},?,?,?)""",
        [(vid, k, *[it.get(f, "") for f in FIELDS], json.dumps(it.get("extra") or {}), _hash(it), it.get("file_dups") or 0)
         for k, it in items.items()])
    chg = [(lob_id, vid, k, "added", None, None, None) for k in added]
    chg += [(lob_id, vid, k, "removed", None, None, None) for k in removed]
    for k, fields in modified.items():
        chg += [(lob_id, vid, k, "modified", f, o, n) for f, o, n in fields]
    c.executemany("""INSERT INTO inventory_changes(lob_id, version_id, item_key, change_type, field, old_value, new_value)
                     VALUES (?,?,?,?,?,?,?)""", chg)

    # rebuild current
    prev_cur = {r["item_key"]: dict(r) for r in c.execute(
        "SELECT item_key, first_version_no, last_changed_version_no FROM inventory_current WHERE lob_id=? AND type_id IS ?",
        (lob_id, type_id))}
    c.execute("DELETE FROM inventory_current WHERE lob_id=? AND type_id IS ?", (lob_id, type_id))
    added_s, mod_s = set(added), set(modified)
    cur_rows = []
    for k, it in items.items():
        p = prev_cur.get(k, {})
        tag = "new" if k in added_s else ("modified" if k in mod_s else "unchanged")
        cur_rows.append((lob_id, k, *[it.get(f, "") for f in FIELDS], json.dumps(it.get("extra") or {}),
                         p.get("first_version_no") or vno,
                         vno if tag != "unchanged" else (p.get("last_changed_version_no") or vno), tag, it.get("file_dups") or 0,
                         type_id))
    c.executemany(
        f"""INSERT INTO inventory_current(lob_id, item_key, {ROW_COLS}, extra, first_version_no, last_changed_version_no, change_tag,
           file_dups, type_id) VALUES (?,?,{ph},?,?,?,?,?,?)""", cur_rows)
    if type_id:
        c.execute("UPDATE lob_types SET current_version_id=? WHERE id=?", (vid, type_id))
    else:
        c.execute("UPDATE lobs SET current_version_id=? WHERE id=?", (vid, lob_id))
    refresh_matches(c, lob_id)
    return {"version_id": vid, "version_no": vno, "type_id": type_id, **{k: v for k, v in summary.items() if k != "samples"}}


def scoped_items(c, lob, parsed, mapping, key_field, scope_msp=None, type_id=None):
    t = get_type(c, lob["id"], type_id)
    items, warnings = build_items(parsed, mapping, key_field, forced_msp=scope_msp, forced_type=t and t["name"],
                                  key_prefix=type_key_prefix(type_id))
    prev = load_version_items(c, stream_version_id(c, lob, type_id))
    if scope_msp:
        items = merge_scope(prev, items, scope_msp)
    return prev, items, warnings


def preview_upload(c, lob_id, token, mapping, key_field, sheet=None, header_row=None, scope_msp=None, type_id=None):
    parsed = parse_upload(token, sheet, header_row)
    lob = db.one(c, "SELECT * FROM lobs WHERE id=?", (lob_id,))
    if not lob:
        raise ValueError("LOB not found")
    prev, items, warnings = scoped_items(c, lob, parsed, mapping, key_field, scope_msp, type_id)
    cur_vid = stream_version_id(c, lob, type_id)
    prev_key = db.one(c, "SELECT key_field FROM inventory_versions WHERE id=?", (cur_vid,)) if cur_vid else None
    if prev_key and prev_key["key_field"] != key_field:
        warnings.insert(0, f"Key field changed from {KEY_FIELDS.get(prev_key['key_field'])} to {KEY_FIELDS.get(key_field)} "
                           "- most rows will show as removed+added.")
    summary, _ = _diff_summary(prev, items)
    summary["warnings"] = warnings
    summary["is_first_version"] = not cur_vid
    summary["type"] = (get_type(c, lob_id, type_id) or {}).get("name")
    summary["scope_msp"] = scope_msp
    summary["file_duplicates"] = sum(it.get("file_dups") or 0 for it in items.values())
    existing = {m["name"].lower() for m in db.rows(c, "SELECT name FROM msps WHERE lob_id=?", (lob_id,))}
    summary["new_msps"] = sorted({(it.get("msp") or "").strip() for it in items.values()
                                  if (it.get("msp") or "").strip() and (it.get("msp") or "").strip().lower() not in existing})
    return summary


def restore_version(c, lob_id, version_id, uploaded_by="", note=""):
    v = db.one(c, "SELECT * FROM inventory_versions WHERE id=? AND lob_id=?", (version_id, lob_id))
    if not v:
        raise ValueError("Version not found")
    items = load_version_items(c, version_id)
    return commit_version(c, lob_id, items, filename=v["filename"], note=note or f"Restored from v{v['version_no']}",
                          uploaded_by=uploaded_by, template_id=v["template_id"], key_field=v["key_field"],
                          mapping=db.jloads(v["mapping"], {}), restored_from=version_id, type_id=v["type_id"])


def compare_versions(c, lob_id, v_old, v_new):
    a = db.one(c, "SELECT id, version_no FROM inventory_versions WHERE id=? AND lob_id=?", (v_old, lob_id))
    b = db.one(c, "SELECT id, version_no FROM inventory_versions WHERE id=? AND lob_id=?", (v_new, lob_id))
    if not a or not b:
        raise ValueError("Version not found")
    summary, _ = _diff_summary(load_version_items(c, a["id"]), load_version_items(c, b["id"]), sample=5000)
    summary["from_version"], summary["to_version"] = a["version_no"], b["version_no"]
    return summary


# ------------------------------------------------------------------ EDR verification
PRESENT = ("Online", "Offline", "Inactive")
INSTALLED_STATES = ("Online", "Offline")
# edr_actual (detailed) -> edr_state (what the dashboards count)
EDR_STATE = {"Online": "Online", "Offline": "Offline", "Inactive": "Offline", "Hidden": "Hidden", "Removed": "Removed",
             "Not Found": "Not Installed", "IP Used by Other Host": "Not Installed"}


def _verify(feasible, claim, actual):
    """Inventory claim vs Falcon reality (data-quality view)."""
    present = actual in PRESENT
    if feasible == "No":
        return "Installed - Marked Not Feasible" if present else "Not Feasible"
    if claim == "Yes":
        if present:
            return "Installed - Inactive" if actual == "Inactive" else "Verified"
        if actual in ("Removed", "Hidden"):
            return "Claimed - Removed from Console"
        return "Claimed - Not Found"
    if claim == "No":
        return "Installed - Marked No" if present else "Pending Install"
    return "Found - No Claim" if present else "Not Found - No Claim"


def is_applicable(live, feasible):
    return (live or "") != "Non Live" and (feasible or "") != "No"


def coverage_status(live, feasible, edr_state):
    if (live or "") == "Non Live":
        return "Non Live"
    if (feasible or "") == "No":
        return "Not Feasible"
    return edr_state


def sync_msps(c, lob_id=None):
    """Create MSP records for MSP names found in inventory and link rows to them."""
    where, params = ("AND lob_id=?", (lob_id,)) if lob_id else ("", ())
    for r in db.rows(c, f"SELECT DISTINCT lob_id, TRIM(msp) msp FROM inventory_current WHERE COALESCE(TRIM(msp),'')<>'' {where}", params):
        c.execute("INSERT OR IGNORE INTO msps(lob_id, name, created_at) VALUES (?,?,?)", (r["lob_id"], r["msp"], db.now_iso()))
    c.execute(f"""UPDATE inventory_current SET msp_id=(SELECT m.id FROM msps m WHERE m.lob_id=inventory_current.lob_id
                  AND m.name=TRIM(inventory_current.msp) COLLATE NOCASE) WHERE 1=1 {where}""", params)


def rebuild_host_map(c):
    c.execute("DELETE FROM host_map")
    c.execute("""INSERT OR IGNORE INTO host_map(aid, lob_id, msp_id, source)
                 SELECT matched_aid, lob_id, MIN(msp_id), 'inventory' FROM inventory_current
                 WHERE matched_aid IS NOT NULL AND edr_state <> 'Not Installed' GROUP BY matched_aid, lob_id""")
    c.execute("""INSERT OR IGNORE INTO host_map(aid, lob_id, msp_id, source)
                 SELECT aid, lob_id, msp_id, 'tag' FROM agent_tags""")


def refresh_matches(c, lob_id=None):
    settings = db.get_settings(c)
    stale_cut = (datetime.now(timezone.utc) - timedelta(days=float(settings.get("inventory_stale_days") or 7))).strftime("%Y-%m-%dT%H:%M:%SZ")
    sync_msps(c, lob_id)
    hosts = {}
    by_ip, by_hn, by_hist = {}, {}, {}
    for r in c.execute("""SELECT aid, hostname, hostname_norm, local_ip, console_state, online_state, last_seen,
                          agent_version, os_version FROM hosts"""):
        hosts[r["aid"]] = r
        if r["local_ip"]:
            by_ip.setdefault(r["local_ip"], []).append(r["aid"])
        if r["hostname_norm"]:
            by_hn.setdefault(r["hostname_norm"], []).append(r["aid"])
    for r in c.execute("SELECT DISTINCT ip, aid FROM ip_history WHERE kind='local'"):
        by_hist.setdefault(r["ip"], set()).add(r["aid"])

    def rank(aid):
        h = hosts[aid]
        return ({"active": 0, "hidden": 1}.get(h["console_state"], 2), -(_ts_num(h["last_seen"])))

    where, params = ("WHERE lob_id=?", (lob_id,)) if lob_id else ("", ())
    items = db.rows(c, f"SELECT lob_id, item_key, ip, node_name, live, edr_feasible, edr_installed FROM inventory_current {where}", params)
    updates = []
    for it in items:
        ip, hn = (it["ip"] or "").strip(), db.norm_hostname(it["node_name"])
        cip = set(by_ip.get(ip, [])) if ip else set()
        chn = set(by_hn.get(hn, [])) if hn else set()
        both = cip & chn
        method, pool = None, set()
        if both:
            method, pool = "ip+hostname", both
        elif chn:
            method, pool = "hostname", chn
        elif cip:
            method, pool = ("ip (hostname differs)" if hn else "ip"), cip
        elif ip and ip in by_hist:
            method, pool = "ip_history", set(by_hist[ip])
        applicable = 1 if is_applicable(it["live"], it["edr_feasible"]) else 0
        if pool:
            best = sorted(pool, key=rank)[0]
            h = hosts[best]
            active_cnt = sum(1 for a in pool if hosts[a]["console_state"] == "active")
            if h["console_state"] == "removed":
                actual = "Removed"
            elif h["console_state"] == "hidden":
                actual = "Hidden"
            elif h["online_state"] == "online":
                actual = "Online"
            elif (h["last_seen"] or "") < stale_cut:
                actual = "Inactive"
            else:
                actual = "Offline"
            weak = method == "ip (hostname differs)" or (method == "ip_history" and hn and h["hostname_norm"] != hn)
            if weak:
                # the IP belongs to a differently named machine: don't count it as installed
                actual = "IP Used by Other Host"
            state = EDR_STATE[actual]
            updates.append((best, method, active_cnt, h["hostname"], h["console_state"], h["online_state"], h["last_seen"],
                            h["agent_version"], h["os_version"], actual,
                            _verify(it["edr_feasible"], it["edr_installed"], "Not Found" if weak else actual),
                            applicable, state, coverage_status(it["live"], it["edr_feasible"], state),
                            it["lob_id"], it["item_key"]))
        else:
            updates.append((None, None, 0, None, None, None, None, None, None, "Not Found",
                            _verify(it["edr_feasible"], it["edr_installed"], "Not Found"),
                            applicable, "Not Installed", coverage_status(it["live"], it["edr_feasible"], "Not Installed"),
                            it["lob_id"], it["item_key"]))
    c.executemany(
        """UPDATE inventory_current SET matched_aid=?, match_method=?, match_count=?, cs_hostname=?, cs_console_state=?,
           cs_online_state=?, cs_last_seen=?, cs_agent_version=?, cs_os=?, edr_actual=?, verification=?,
           applicable=?, edr_state=?, coverage_status=?
           WHERE lob_id=? AND item_key=?""", updates)
    tag_duplicates(c, lob_id)
    rebuild_host_map(c)


def tag_duplicates(c, lob_id=None):
    """Count rows per LOB sharing an IP / node name. dup_ip or dup_name > 1 marks a duplicate."""
    where, params = ("WHERE lob_id=?", (lob_id,)) if lob_id else ("", ())
    rows = db.rows(c, f"SELECT lob_id, item_key, ip, node_name FROM inventory_current {where}", params)
    by_ip, by_name = {}, {}
    for r in rows:
        ip, nn = (r["ip"] or "").strip().lower(), db.norm_hostname(r["node_name"])
        r["_ip"], r["_nn"] = ip, nn
        if ip:
            by_ip[(r["lob_id"], ip)] = by_ip.get((r["lob_id"], ip), 0) + 1
        if nn:
            by_name[(r["lob_id"], nn)] = by_name.get((r["lob_id"], nn), 0) + 1
    c.executemany("UPDATE inventory_current SET dup_ip=?, dup_name=? WHERE lob_id=? AND item_key=?",
                  [(by_ip.get((r["lob_id"], r["_ip"]), 0), by_name.get((r["lob_id"], r["_nn"]), 0), r["lob_id"], r["item_key"])
                   for r in rows])


def _ts_num(ts):
    return int(re.sub(r"\D", "", ts or "") or 0)


# ------------------------------------------------------------------ agent tagging (AID + MSP sheet)
def resolve_agents(c, values):
    """Map AID / hostname / IP values to active agent IDs. Returns {value: aid or None}."""
    vals = [v.strip() for v in values if v and v.strip()]
    out = {v: None for v in vals}
    for i in range(0, len(vals), 500):
        ch = vals[i:i + 500]
        ph = ",".join("?" * len(ch))
        low = [v.lower() for v in ch]
        for r in c.execute(f"SELECT aid FROM hosts WHERE aid IN ({ph})", low):
            for v in ch:
                if v.lower() == r["aid"]:
                    out[v] = r["aid"]
        norm = {db.norm_hostname(v): v for v in ch if out[v] is None}
        if norm:
            ph2 = ",".join("?" * len(norm))
            for r in c.execute(f"""SELECT aid, hostname_norm FROM hosts WHERE hostname_norm IN ({ph2})
                                   ORDER BY console_state='active' DESC, last_seen DESC""", list(norm)):
                v = norm[r["hostname_norm"]]
                if out[v] is None:
                    out[v] = r["aid"]
        ipv = [v for v in ch if out[v] is None and db.ip_to_num(v) is not None]
        if ipv:
            ph3 = ",".join("?" * len(ipv))
            for r in c.execute(f"""SELECT aid, local_ip FROM hosts WHERE local_ip IN ({ph3})
                                   ORDER BY console_state='active' DESC, last_seen DESC""", ipv):
                if out.get(r["local_ip"]) is None:
                    out[r["local_ip"]] = r["aid"]
    return out


def tag_rows(c, lob_id, parsed, id_col, msp_col=None, msp_id=None):
    """Returns (rows, summary) for an agent-tag sheet without writing anything."""
    idx = {h: i for i, h in enumerate(parsed["headers"])}
    if id_col not in idx:
        raise ValueError("Pick the column that holds the Agent ID (or hostname / IP)")
    msps = {m["name"].lower(): m for m in db.rows(c, "SELECT id, name FROM msps WHERE lob_id=?", (lob_id,))}
    fixed = db.one(c, "SELECT id, name FROM msps WHERE id=? AND lob_id=?", (msp_id, lob_id)) if msp_id else None
    raw = []
    for r in parsed["rows"]:
        val = r[idx[id_col]].strip()
        if not val:
            continue
        msp = fixed["name"] if fixed else (r[idx[msp_col]].strip() if msp_col in idx else "")
        raw.append((val, msp))
    resolved = resolve_agents(c, [v for v, _ in raw])
    inv = {r["matched_aid"] for r in c.execute(
        "SELECT matched_aid FROM inventory_current WHERE lob_id=? AND matched_aid IS NOT NULL AND edr_state<>'Not Installed'", (lob_id,))}
    rows, new_msps = [], set()
    for val, msp in raw:
        aid = resolved.get(val)
        if msp and msp.lower() not in msps:
            new_msps.add(msp)
        rows.append({"value": val, "aid": aid, "msp": msp, "status": "Not found in Falcon" if not aid else ("Already in inventory" if aid in inv else "Not in inventory")})
    summary = {"total": len(rows), "resolved": sum(1 for r in rows if r["aid"]), "unresolved": sum(1 for r in rows if not r["aid"]),
               "in_inventory": sum(1 for r in rows if r["status"] == "Already in inventory"),
               "unlisted": sum(1 for r in rows if r["status"] == "Not in inventory"), "new_msps": sorted(new_msps)}
    return rows, summary


def commit_tags(c, lob_id, rows, source, replace_scope=None):
    """replace_scope: None = add/update, 'lob' = replace all tags of the LOB, int msp_id = replace that MSP's tags."""
    if replace_scope == "lob":
        c.execute("DELETE FROM agent_tags WHERE lob_id=?", (lob_id,))
    elif replace_scope:
        c.execute("DELETE FROM agent_tags WHERE lob_id=? AND msp_id=?", (lob_id, replace_scope))
    for r in rows:
        if r["msp"]:
            c.execute("INSERT OR IGNORE INTO msps(lob_id, name, created_at) VALUES (?,?,?)", (lob_id, r["msp"], db.now_iso()))
    msps = {m["name"].lower(): m["id"] for m in db.rows(c, "SELECT id, name FROM msps WHERE lob_id=?", (lob_id,))}
    now = db.now_iso()
    data = [(r["aid"], lob_id, msps.get((r["msp"] or "").lower()), source, now) for r in rows if r["aid"]]
    c.executemany("INSERT OR REPLACE INTO agent_tags(aid, lob_id, msp_id, source, tagged_at) VALUES (?,?,?,?,?)", data)
    rebuild_host_map(c)
    return len(data)
