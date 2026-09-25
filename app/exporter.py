import io
import re
from datetime import datetime

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")


def _clean(v):
    if isinstance(v, str):
        return _ILLEGAL.sub("", v)[:32000]
    return v


def xlsx_response(sheets, filename):
    """sheets: list of (title, columns[(key,label)], rows[dict])"""
    wb = Workbook(write_only=True)
    for title, columns, rows in sheets:
        ws = wb.create_sheet(title=re.sub(r"[\[\]\*\?/\\:]", "_", title)[:31] or "Sheet")
        widths = [max(10, min(45, len(lbl) + 2)) for _, lbl in columns]
        for i, rw in enumerate(rows[:200]):
            for j, (k, _) in enumerate(columns):
                widths[j] = max(widths[j], min(45, len(str(rw.get(k) or "")) + 2))
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w
        ws.freeze_panes = "A2"
        head = []
        for _, lbl in columns:
            cell = WriteOnlyCell(ws, value=lbl)
            cell.font, cell.fill = HEAD_FONT, HEAD_FILL
            cell.alignment = Alignment(vertical="center")
            head.append(cell)
        ws.append(head)
        for r in rows:
            ws.append([_clean(r.get(k)) for k, _ in columns])
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    fname = f"{filename}_{stamp}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
